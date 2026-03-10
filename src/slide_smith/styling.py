from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Pt


@dataclass(frozen=True)
class TextStyle:
    # Run-level
    font: str | None = None
    size_pt: int | float | None = None
    color_hex: str | None = None  # e.g. "112233"
    bold: bool | None = None
    italic: bool | None = None

    # Paragraph-level
    align: str | None = None  # 'left'|'center'|'right'|'justify'
    line_spacing_pt: int | float | None = None


_ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}


def _parse_style(obj: dict[str, Any] | None) -> TextStyle:
    if not obj:
        return TextStyle()
    return TextStyle(
        font=obj.get("font"),
        size_pt=obj.get("size_pt"),
        color_hex=obj.get("color_hex"),
        bold=obj.get("bold"),
        italic=obj.get("italic"),
        align=obj.get("align"),
        line_spacing_pt=obj.get("line_spacing_pt"),
    )


def load_styles(template_spec: dict[str, Any]) -> dict[str, TextStyle]:
    styles = template_spec.get("styles") or {}
    return {
        "title": _parse_style(styles.get("title")),
        "subtitle": _parse_style(styles.get("subtitle")),
        "body": _parse_style(styles.get("body")),
        "bullets": _parse_style(styles.get("bullets")),
    }


def apply_text_style(text_frame, style: TextStyle) -> None:
    """Best-effort style application across all paragraphs/runs."""
    for paragraph in text_frame.paragraphs:
        # Paragraph-level
        if style.align and style.align in _ALIGN:
            paragraph.alignment = _ALIGN[style.align]
        if style.line_spacing_pt is not None:
            paragraph.line_spacing = Pt(style.line_spacing_pt)

        # Ensure at least one run exists in common cases.
        runs = list(paragraph.runs)
        if not runs and paragraph.text:
            _ = paragraph.add_run()
            runs = list(paragraph.runs)

        for run in runs:
            font = run.font
            if style.font:
                font.name = style.font
            if style.size_pt is not None:
                font.size = Pt(style.size_pt)
            if style.color_hex:
                font.color.rgb = RGBColor.from_string(style.color_hex.upper())
            if style.bold is not None:
                font.bold = style.bold
            if style.italic is not None:
                font.italic = style.italic
